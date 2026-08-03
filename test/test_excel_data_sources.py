"""Offline tests for converting Excel workbooks into SQLite data sources."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from infrastructure import data_sources


class ExcelDataSourceTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.upload_directory = Path(self.temporary_directory.name).resolve()
        self.directory_patch = patch.object(
            data_sources,
            "UPLOAD_DIRECTORY",
            self.upload_directory,
        )
        self.directory_patch.start()

    def tearDown(self) -> None:
        self.directory_patch.stop()
        self.temporary_directory.cleanup()

    @staticmethod
    def _workbook_bytes() -> bytes:
        workbook = Workbook()
        sales = workbook.active
        sales.title = "销售 数据"
        sales.append([])
        sales.append(["订单 ID", "金额", "日期", "金额", None])
        sales.append([1, 12.5, datetime(2026, 8, 1, 9, 30), "含税", "华东"])
        sales.append([2, 8, datetime(2026, 8, 2, 10, 0), "未税", "华南"])
        sales.append([None, None, None, None, None])

        customers = workbook.create_sheet("客户")
        customers.append(["名称", "活跃"])
        customers.append(["甲公司", True])

        empty = workbook.create_sheet("空白")
        empty.append([])

        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    def test_excel_workbook_becomes_query_ready_sqlite_snapshot(self) -> None:
        snapshot = data_sources.stage_excel_workbook(
            "八月经营数据.xlsx",
            self._workbook_bytes(),
        )
        self.addCleanup(snapshot.staged_path.unlink, missing_ok=True)

        self.assertTrue(snapshot.staged_path.is_file())
        self.assertEqual(snapshot.target_path.parent, self.upload_directory)
        self.assertEqual(snapshot.target_path.suffix, ".db")
        self.assertEqual(snapshot.display_name, "八月经营数据")
        self.assertEqual(snapshot.staged_path.stat().st_mode & 0o777, 0o600)

        catalog = data_sources.inspect_sqlite_database(snapshot.staged_path)
        self.assertEqual(
            [table["name"] for table in catalog["tables"]],
            ["客户", "销售_数据"],
        )
        sales = next(
            table for table in catalog["tables"] if table["name"] == "销售_数据"
        )
        self.assertEqual(
            [column["name"] for column in sales["columns"]],
            ["订单_ID", "金额", "日期", "金额_2", "column_5"],
        )
        self.assertEqual(
            [column["type"] for column in sales["columns"]],
            ["INTEGER", "REAL", "TEXT", "TEXT", "TEXT"],
        )

        with sqlite3.connect(snapshot.staged_path) as connection:
            rows = connection.execute(
                'SELECT "订单_ID", "金额", "column_5" '
                'FROM "销售_数据" ORDER BY "订单_ID"'
            ).fetchall()
            metadata = dict(
                connection.execute(
                    "SELECT key, value FROM _chatbi_source_metadata"
                ).fetchall()
            )
        self.assertEqual(rows, [(1, 12.5, "华东"), (2, 8.0, "华南")])
        self.assertEqual(metadata["database_type"], "Excel")
        self.assertEqual(metadata["source_file_name"], "八月经营数据.xlsx")
        self.assertEqual(metadata["sheet_count"], "2")
        self.assertEqual(metadata["row_count"], "3")

    def test_invalid_excel_upload_is_rejected_without_leftovers(self) -> None:
        with self.assertRaisesRegex(
            data_sources.DataSourceError,
            "只支持 .xlsx 和 .xlsm",
        ):
            data_sources.stage_excel_workbook("legacy.xls", b"not-an-excel")
        with self.assertRaisesRegex(
            data_sources.DataSourceError,
            "不是有效的 Excel",
        ):
            data_sources.stage_excel_workbook("broken.xlsx", b"not-an-excel")
        with patch.object(data_sources, "MAX_EXCEL_COLUMNS", 4):
            with self.assertRaisesRegex(
                data_sources.DataSourceError,
                "超过 4 列限制",
            ):
                data_sources.stage_excel_workbook(
                    "too-wide.xlsx",
                    self._workbook_bytes(),
                )
        self.assertEqual(list(self.upload_directory.iterdir()), [])


if __name__ == "__main__":
    unittest.main(verbosity=2)
